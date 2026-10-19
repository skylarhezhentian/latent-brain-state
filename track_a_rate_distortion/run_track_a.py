"""Execute Track A: rate-distortion of low-dimensional spatial LFP representations.

The experiment isolates representation quality:
  1. Start from the same destriped 384-channel LFP at 250 Hz.
  2. Fit a spatial PCA basis on the first 80% of each probe's time series.
  3. Reconstruct the held-out final 20% at k = 2, 4, 8, 16, 32, 64 dimensions.
  4. Measure native waveform reconstruction R2.
  5. Recompute the same depth-group band-power features from raw and reconstructed LFP,
     and use the same ridge-decoding procedure to quantify behavior preservation.

Outputs include CSV/JSON results, figures, and a presentation-ready PowerPoint deck.
"""
from __future__ import annotations

import json
import math
import os
import time
import traceback
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Sequence, Tuple

os.environ.setdefault("OMP_NUM_THREADS", "1")
os.environ.setdefault("MKL_NUM_THREADS", "1")
os.environ.setdefault("OPENBLAS_NUM_THREADS", "1")
os.environ.setdefault("NUMEXPR_NUM_THREADS", "1")

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import scipy.signal
from sklearn.decomposition import PCA
from sklearn.linear_model import RidgeCV
from sklearn.metrics import r2_score
from sklearn.model_selection import TimeSeriesSplit
from sklearn.preprocessing import StandardScaler

from one.api import ONE
from brainbox.io.one import SessionLoader, SpikeSortingLoader
from ibldsp.voltage import destripe_lfp


# Three public BWM insertions chosen before running this analysis.
# The first two are simultaneous probes from one session; the third is from another lab/session.
PROBES = [
    {
        "pid": "63517fd4-ece1-49eb-9259-371dc30b1dd6",
        "eid": "4a45c8ba-db6f-4f11-9403-56e06a33dfa4",
        "t0_main_s": 50.01006565988064,
        "label": "Dan Lab · probe00",
    },
    {
        "pid": "ca5404f7-297c-40f1-bbf0-5ac0a63e24f8",
        "eid": "4a45c8ba-db6f-4f11-9403-56e06a33dfa4",
        "t0_main_s": 50.01006565988064,
        "label": "Dan Lab · probe01",
    },
    {
        "pid": "db2100c3-64ac-463e-97a1-20ce8266cd5f",
        "eid": "0cc486c3-8c7b-494d-aa04-b70e2690bcba",
        "t0_main_s": 56.682618141174316,
        "label": "Witten Lab · probe00",
    },
]

OUT = Path("track_a_outputs")
CACHE = OUT / "cache"
FIG = OUT / "figures"
DURATION_S = 60.0
FS_OUT = 250.0
BEHAVIOR_STEP_S = 0.1
DIMS = (2, 4, 8, 16, 32, 64)
TRAIN_FRACTION = 0.80
BANDS = {
    "delta": (1.0, 4.0),
    "theta": (4.0, 8.0),
    "alpha": (8.0, 12.0),
    "beta": (12.0, 30.0),
    "gamma": (30.0, 80.0),
}
COLORS = {
    "ink": "#0B0B0B",
    "gray": "#5B5A57",
    "grid": "#E6E5E1",
    "blue": "#2A78D6",
    "orange": "#EB6834",
    "green": "#1BAF7A",
    "purple": "#7A58C2",
    "red": "#C94F5C",
}


@dataclass
class ProbeResult:
    pid: str
    eid: str
    label: str
    k: int
    values_per_second: float
    channel_reduction: float
    reconstruction_r2: float
    reconstruction_rmse_uv: float
    behavior_r2_mean: float
    behavior_delta_r2_mean: float
    n_behaviors: int
    behavior_scores: Dict[str, float]
    raw_behavior_scores: Dict[str, float]


def _first_existing(df: pd.DataFrame, candidates: Sequence[str]) -> str | None:
    return next((c for c in candidates if c in df.columns), None)


def _read_stream(sr, first: int, last: int, piece_s: float = 5.0) -> np.ndarray:
    """Read [channels, samples] float32 data while avoiding fragile huge range requests."""
    step = max(1, int(round(piece_s * float(sr.fs))))
    pieces = []
    for a in range(first, last, step):
        b = min(a + step, last)
        pieces.append(np.asarray(sr[a:b, :-sr.nsync], dtype=np.float32))
    if not pieces:
        raise RuntimeError("No samples returned from LFP stream")
    return np.concatenate(pieces, axis=0).T


def load_destriped_250(one: ONE, pid: str, t0_main_s: float, duration_s: float) -> Tuple[np.ndarray, float]:
    """Stream, destripe at native rate, then anti-alias and resample to 250 Hz."""
    CACHE.mkdir(parents=True, exist_ok=True)
    cache_file = CACHE / f"{pid}_destriped250.npz"
    if cache_file.exists():
        with np.load(cache_file) as z:
            return z["x_uv"].astype(np.float32), float(z["fs"])

    ssl = SpikeSortingLoader(pid=pid, one=one)
    sr = ssl.raw_electrophysiology(band="lf", stream=True)
    fs = float(sr.fs)
    if not np.isclose(fs / FS_OUT, round(fs / FS_OUT), rtol=0, atol=1e-6):
        raise ValueError(f"Native fs={fs} is not an integer multiple of {FS_OUT}")
    down = int(round(fs / FS_OUT))

    try:
        t0_probe = float(ssl.timesprobe2times(np.array([t0_main_s]), direction="reverse")[0])
    except Exception:
        t0_probe = float(t0_main_s)

    s0 = int(round(t0_probe * fs))
    n_total = int(round(duration_s * fs))
    chunk = int(round(10.0 * fs))
    pad = int(round(1.0 * fs))
    pieces: List[np.ndarray] = []

    try:
        for rel0 in range(0, n_total, chunk):
            rel1 = min(rel0 + chunk, n_total)
            desired0 = s0 + rel0
            desired1 = s0 + rel1
            a = max(0, desired0 - pad)
            b = desired1 + pad
            raw = _read_stream(sr, a, b)
            x = destripe_lfp(
                raw,
                fs=fs,
                channel_labels=np.zeros(raw.shape[0]),
            ).astype(np.float32)
            y = scipy.signal.resample_poly(x, up=1, down=down, axis=1).astype(np.float32)
            keep0 = int(round((desired0 - a) / down))
            keep_n = int(round((desired1 - desired0) / down))
            pieces.append(y[:, keep0:keep0 + keep_n])
    finally:
        try:
            sr.close()
        except Exception:
            pass

    x_uv = np.concatenate(pieces, axis=1) * 1e6
    target_n = int(round(duration_s * FS_OUT))
    x_uv = x_uv[:, :target_n].astype(np.float32)
    if x_uv.shape[0] != 384 or x_uv.shape[1] < target_n - 2:
        raise RuntimeError(f"Unexpected LFP shape {x_uv.shape}; expected approximately (384, {target_n})")
    np.savez_compressed(cache_file, x_uv=x_uv, fs=FS_OUT)
    return x_uv, FS_OUT


def load_behavior_series(one: ONE, eid: str) -> Dict[str, Tuple[np.ndarray, np.ndarray]]:
    """Load continuous wheel and motion-energy behavior from one session."""
    sl = SessionLoader(one=one, eid=eid)
    sl.load_session_data(
        trials=False,
        wheel=True,
        pose=False,
        motion_energy=True,
        pupil=False,
    )
    series: Dict[str, Tuple[np.ndarray, np.ndarray]] = {}

    wheel = getattr(sl, "wheel", None)
    if isinstance(wheel, pd.DataFrame) and len(wheel):
        tc = _first_existing(wheel, ("times", "timestamps"))
        vc = _first_existing(wheel, ("velocity", "speed"))
        if tc and vc:
            series["wheel_speed"] = (
                wheel[tc].to_numpy(dtype=float),
                np.abs(wheel[vc].to_numpy(dtype=float)),
            )

    motion = getattr(sl, "motion_energy", None)
    if isinstance(motion, Mapping):
        specs = (
            ("leftCamera", ("whiskerMotionEnergy", "motionEnergy"), "whisker_left"),
            ("rightCamera", ("whiskerMotionEnergy", "motionEnergy"), "whisker_right"),
            ("bodyCamera", ("bodyMotionEnergy", "motionEnergy"), "body_motion"),
        )
        for camera, value_candidates, output_name in specs:
            df = motion.get(camera)
            if not isinstance(df, pd.DataFrame) or not len(df):
                continue
            tc = _first_existing(df, ("times", "timestamps"))
            vc = _first_existing(df, value_candidates)
            if tc and vc:
                series[output_name] = (
                    df[tc].to_numpy(dtype=float),
                    df[vc].to_numpy(dtype=float),
                )
    return series


def behavior_at_times(
    series: Mapping[str, Tuple[np.ndarray, np.ndarray]],
    times_abs: np.ndarray,
) -> Tuple[np.ndarray, List[str]]:
    names = sorted(series)
    y = np.full((len(times_abs), len(names)), np.nan, dtype=np.float64)
    for j, name in enumerate(names):
        t, v = series[name]
        good = np.isfinite(t) & np.isfinite(v)
        if good.sum() < 10:
            continue
        t = t[good]
        v = v[good]
        order = np.argsort(t)
        t, v = t[order], v[order]
        y[:, j] = np.interp(times_abs, t, v, left=np.nan, right=np.nan)
    return y, names


def group_channels(x_uv: np.ndarray, groups: int = 24) -> np.ndarray:
    if x_uv.shape[0] % groups:
        raise ValueError(f"{x_uv.shape[0]} channels cannot be split into {groups} equal groups")
    per = x_uv.shape[0] // groups
    return x_uv.reshape(groups, per, x_uv.shape[1]).mean(axis=1)


def bandpower_features(x_uv: np.ndarray, fs: float) -> Tuple[np.ndarray, np.ndarray, List[str]]:
    """Fixed downstream feature extractor used for raw and reconstructed LFP."""
    grouped = group_channels(x_uv, 24)
    nperseg = int(round(1.0 * fs))
    step = int(round(BEHAVIOR_STEP_S * fs))
    noverlap = nperseg - step
    f, t, psd = scipy.signal.spectrogram(
        grouped,
        fs=fs,
        window="hann",
        nperseg=nperseg,
        noverlap=noverlap,
        nfft=max(256, nperseg),
        detrend="constant",
        scaling="density",
        mode="psd",
        axis=-1,
    )
    feats: List[np.ndarray] = []
    names: List[str] = []
    for band, (lo, hi) in BANDS.items():
        mask = (f >= lo) & (f < hi)
        if not np.any(mask):
            continue
        p = np.log10(np.mean(psd[:, mask, :], axis=1) + 1e-12)
        feats.append(p.T)
        names.extend([f"depth{g:02d}_{band}" for g in range(p.shape[0])])
    X = np.concatenate(feats, axis=1)
    return X.astype(np.float64), t.astype(np.float64), names


def pooled_reconstruction_metrics(x: np.ndarray, xhat: np.ndarray) -> Tuple[float, float]:
    err = np.sum((x - xhat) ** 2, dtype=np.float64)
    centered = x - np.mean(x, axis=1, keepdims=True)
    den = np.sum(centered ** 2, dtype=np.float64)
    r2 = float(1.0 - err / den) if den > 0 else float("nan")
    rmse = float(np.sqrt(np.mean((x - xhat) ** 2, dtype=np.float64)))
    return r2, rmse


def decode_behaviors(
    features: np.ndarray,
    t_rel: np.ndarray,
    y: np.ndarray,
    behavior_names: Sequence[str],
    duration_s: float,
) -> Dict[str, float]:
    """Identical ridge procedure for raw and every reconstructed condition."""
    train_limit = TRAIN_FRACTION * duration_s
    # A one-second spectral window is centered at each t_rel. Keep windows wholly inside a split.
    train_time = t_rel <= (train_limit - 0.5)
    test_time = t_rel >= (train_limit + 0.5)
    scores: Dict[str, float] = {}
    alphas = np.logspace(-2, 4, 13)

    for j, name in enumerate(behavior_names):
        valid = np.isfinite(y[:, j]) & np.all(np.isfinite(features), axis=1)
        tr = valid & train_time
        te = valid & test_time
        if tr.sum() < 80 or te.sum() < 30 or np.nanstd(y[te, j]) < 1e-10:
            scores[name] = float("nan")
            continue
        scaler = StandardScaler()
        Xtr = scaler.fit_transform(features[tr])
        Xte = scaler.transform(features[te])
        n_splits = min(5, max(2, tr.sum() // 50))
        model = RidgeCV(
            alphas=alphas,
            cv=TimeSeriesSplit(n_splits=n_splits),
            scoring="neg_mean_squared_error",
        )
        model.fit(Xtr, y[tr, j])
        pred = model.predict(Xte)
        scores[name] = float(r2_score(y[te, j], pred))
    return scores


def finite_mean(values: Iterable[float]) -> float:
    x = np.asarray(list(values), dtype=float)
    return float(np.nanmean(x)) if np.isfinite(x).any() else float("nan")


def run_one_probe(one: ONE, spec: Mapping[str, object]) -> Tuple[List[ProbeResult], Dict[str, object]]:
    pid, eid = str(spec["pid"]), str(spec["eid"])
    label = str(spec["label"])
    t0 = float(spec["t0_main_s"])
    print(f"\n[{pid[:8]}] loading {label}", flush=True)
    x_uv, fs = load_destriped_250(one, pid, t0, DURATION_S)
    behavior_series = load_behavior_series(one, eid)
    if not behavior_series:
        raise RuntimeError("No continuous behavior variables were available")

    n = x_uv.shape[1]
    train_n = int(math.floor(TRAIN_FRACTION * n))
    pca = PCA(n_components=max(DIMS), svd_solver="randomized", random_state=0)
    pca.fit(x_uv[:, :train_n].T)
    all_scores = pca.transform(x_uv.T)

    raw_features, feature_times, _ = bandpower_features(x_uv, fs)
    y, behavior_names = behavior_at_times(behavior_series, t0 + feature_times)
    raw_behavior = decode_behaviors(raw_features, feature_times, y, behavior_names, DURATION_S)
    raw_mean = finite_mean(raw_behavior.values())

    rows: List[ProbeResult] = []
    recon_examples: Dict[int, np.ndarray] = {}
    x_test = x_uv[:, train_n:]

    for k in DIMS:
        comp = pca.components_[:k]
        xhat = (all_scores[:, :k] @ comp + pca.mean_).T.astype(np.float32)
        r2, rmse = pooled_reconstruction_metrics(x_test, xhat[:, train_n:])
        features, times, _ = bandpower_features(xhat, fs)
        behavior = decode_behaviors(features, times, y, behavior_names, DURATION_S)
        mean_r2 = finite_mean(behavior.values())
        rows.append(
            ProbeResult(
                pid=pid,
                eid=eid,
                label=label,
                k=int(k),
                values_per_second=float(k * fs),
                channel_reduction=float(384 / k),
                reconstruction_r2=r2,
                reconstruction_rmse_uv=rmse,
                behavior_r2_mean=mean_r2,
                behavior_delta_r2_mean=float(mean_r2 - raw_mean),
                n_behaviors=int(sum(np.isfinite(list(behavior.values())))),
                behavior_scores=behavior,
                raw_behavior_scores=raw_behavior,
            )
        )
        if k in (4, 16, 64):
            recon_examples[k] = xhat
        print(
            f"[{pid[:8]}] k={k:>2} | recon R2={r2:.4f} | "
            f"behavior R2={mean_r2:.4f} | delta={mean_r2 - raw_mean:+.4f}",
            flush=True,
        )

    example = {
        "pid": pid,
        "label": label,
        "x_uv": x_uv,
        "recon": recon_examples,
        "fs": fs,
        "raw_behavior": raw_behavior,
        "behavior_names": behavior_names,
        "explained_variance_ratio": pca.explained_variance_ratio_,
    }
    return rows, example


def _style_axis(ax):
    ax.set_facecolor("white")
    ax.grid(True, color=COLORS["grid"], linewidth=0.8)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(colors=COLORS["gray"], labelsize=10)


def aggregate_results(df: pd.DataFrame) -> pd.DataFrame:
    numeric = [
        "values_per_second",
        "channel_reduction",
        "reconstruction_r2",
        "reconstruction_rmse_uv",
        "behavior_r2_mean",
        "behavior_delta_r2_mean",
    ]
    agg = df.groupby("k")[numeric].median().reset_index()
    agg["n_probes"] = df.groupby("k").pid.nunique().reindex(agg.k).to_numpy()
    return agg


def make_figures(df: pd.DataFrame, examples: Sequence[Dict[str, object]]) -> Dict[str, str]:
    FIG.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11})
    agg = aggregate_results(df)
    paths: Dict[str, str] = {}

    # 1) Reconstruction rate-distortion curve.
    fig, ax = plt.subplots(figsize=(10.8, 5.8), facecolor="white")
    _style_axis(ax)
    for _, g in df.groupby("pid"):
        g = g.sort_values("k")
        ax.plot(g.k, g.reconstruction_r2, color=COLORS["blue"], alpha=0.20, linewidth=1.2)
    ax.plot(agg.k, agg.reconstruction_r2, marker="o", markersize=7,
            color=COLORS["blue"], linewidth=2.8, label="median across probes")
    for _, r in agg.iterrows():
        ax.annotate(f"{r.reconstruction_r2:.3f}", (r.k, r.reconstruction_r2),
                    xytext=(0, 9), textcoords="offset points", ha="center", fontsize=9)
    ax.set_xscale("log", base=2)
    ax.set_xticks(list(DIMS), [str(k) for k in DIMS])
    ax.set_xlabel("representation dimension k  (spatial PCA components)")
    ax.set_ylabel("held-out reconstruction R²")
    ax.set_ylim(max(-0.05, float(agg.reconstruction_r2.min()) - 0.05), 1.01)
    ax.set_title("Rate–distortion: fewer dimensions versus waveform fidelity", loc="left",
                 fontsize=17, fontweight="bold", color=COLORS["ink"])
    ax.legend(frameon=False, loc="lower right")
    fig.tight_layout()
    p = FIG / "rate_distortion_reconstruction.png"
    fig.savefig(p, dpi=220, bbox_inches="tight")
    plt.close(fig)
    paths["reconstruction"] = str(p)

    # 2) Behavior loss relative to raw, per behavior.
    behavior_names = sorted({name for scores in df.behavior_scores for name in scores})
    fig, ax = plt.subplots(figsize=(10.8, 5.8), facecolor="white")
    _style_axis(ax)
    palette = [COLORS["blue"], COLORS["orange"], COLORS["green"], COLORS["purple"], COLORS["red"]]
    any_line = False
    for color, name in zip(palette, behavior_names):
        rows = []
        for _, r in df.iterrows():
            score = r.behavior_scores.get(name, np.nan)
            raw = r.raw_behavior_scores.get(name, np.nan)
            rows.append({"k": r.k, "pid": r.pid, "delta": score - raw})
        bd = pd.DataFrame(rows)
        med = bd.groupby("k").delta.median().reindex(DIMS)
        if np.isfinite(med).any():
            ax.plot(DIMS, med, marker="o", linewidth=2.2, color=color, label=name.replace("_", " "))
            any_line = True
    ax.axhline(0, color=COLORS["ink"], linewidth=1.0)
    ax.set_xscale("log", base=2)
    ax.set_xticks(list(DIMS), [str(k) for k in DIMS])
    ax.set_xlabel("representation dimension k")
    ax.set_ylabel("behavior R² change relative to raw LFP")
    ax.set_title("Behavioral information is tested independently of waveform reconstruction", loc="left",
                 fontsize=16, fontweight="bold", color=COLORS["ink"])
    if any_line:
        ax.legend(frameon=False, ncol=2, loc="lower right")
    fig.tight_layout()
    p = FIG / "behavior_preservation.png"
    fig.savefig(p, dpi=220, bbox_inches="tight")
    plt.close(fig)
    paths["behavior"] = str(p)

    # 3) Combined rate versus both forms of loss.
    fig, ax1 = plt.subplots(figsize=(10.8, 5.8), facecolor="white")
    _style_axis(ax1)
    ax1.plot(agg.k, 1 - agg.reconstruction_r2, marker="o", linewidth=2.6,
             color=COLORS["blue"], label="waveform distortion: 1 − R²")
    ax1.set_xscale("log", base=2)
    ax1.set_yscale("log")
    ax1.set_xticks(list(DIMS), [str(k) for k in DIMS])
    ax1.set_xlabel("representation dimension k")
    ax1.set_ylabel("waveform distortion (log scale)", color=COLORS["blue"])
    ax1.tick_params(axis="y", colors=COLORS["blue"])
    ax2 = ax1.twinx()
    behavior_loss = np.maximum(0, -agg.behavior_delta_r2_mean.to_numpy())
    ax2.plot(agg.k, behavior_loss, marker="s", linewidth=2.6,
             color=COLORS["orange"], label="mean behavior R² loss")
    ax2.set_ylabel("behavior R² loss relative to raw", color=COLORS["orange"])
    ax2.tick_params(axis="y", colors=COLORS["orange"])
    ax1.set_title("The two distortions do not have to flatten at the same rate", loc="left",
                  fontsize=16, fontweight="bold", color=COLORS["ink"])
    lines = ax1.get_lines() + ax2.get_lines()
    ax1.legend(lines, [line.get_label() for line in lines], frameon=False, loc="upper right")
    fig.tight_layout()
    p = FIG / "two_distortions.png"
    fig.savefig(p, dpi=220, bbox_inches="tight")
    plt.close(fig)
    paths["combined"] = str(p)

    # 4) Representative traces.
    ex = examples[0]
    raw = np.asarray(ex["x_uv"])
    recon = ex["recon"]
    fs = float(ex["fs"])
    start = int(round((TRAIN_FRACTION * DURATION_S + 2.0) * fs))
    length = int(round(2.0 * fs))
    channels = (64, 192, 320)
    fig, axes = plt.subplots(len(channels), 1, figsize=(11, 6.3), sharex=True, facecolor="white")
    tt = np.arange(length) / fs
    offsets = {4: 150.0, 16: 300.0, 64: 450.0}
    for ax, ch in zip(axes, channels):
        _style_axis(ax)
        seg = raw[ch, start:start + length]
        seg = seg - np.mean(seg)
        ax.plot(tt, seg, color=COLORS["ink"], linewidth=1.2, label="raw")
        for k, color in zip((4, 16, 64), (COLORS["red"], COLORS["orange"], COLORS["blue"])):
            y = np.asarray(recon[k])[ch, start:start + length]
            y = y - np.mean(y) + offsets[k]
            ax.plot(tt, y, color=color, linewidth=1.0, label=f"k={k}" if ch == channels[0] else None)
        ax.set_ylabel(f"ch {ch}\nµV + offset", fontsize=9)
    axes[0].legend(frameon=False, ncol=4, loc="upper right")
    axes[-1].set_xlabel("time in held-out interval (s)")
    fig.suptitle("What increasing representation dimension restores", x=0.08, ha="left",
                 fontsize=16, fontweight="bold", color=COLORS["ink"])
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    p = FIG / "example_reconstruction.png"
    fig.savefig(p, dpi=220, bbox_inches="tight")
    plt.close(fig)
    paths["traces"] = str(p)
    return paths


def _add_text(slide, text, x, y, w, h, size=20, bold=False, color="#0B0B0B",
              font="Arial", align=None):
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = _rgb(color)
    if align:
        p.alignment = getattr(PP_ALIGN, align)
    return box


def _rgb(hex_color: str):
    from pptx.dml.color import RGBColor
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _add_title(slide, title: str, kicker: str | None = None):
    if kicker:
        _add_text(slide, kicker.upper(), 0.65, 0.32, 12.0, 0.25, size=10, bold=True, color=COLORS["blue"])
    _add_text(slide, title, 0.65, 0.62, 12.0, 0.65, size=28, bold=True, color=COLORS["ink"])


def _add_footer(slide, text: str):
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(7.18), Inches(12.0), Inches(0.012))
    line.fill.solid(); line.fill.fore_color.rgb = _rgb(COLORS["grid"]); line.line.fill.background()
    _add_text(slide, text, 0.65, 7.24, 12.0, 0.18, size=8, color=COLORS["gray"])


def make_deck(df: pd.DataFrame, figure_paths: Mapping[str, str], failures: Sequence[Mapping[str, str]]) -> Path:
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE

    agg = aggregate_results(df)
    n_probes = int(df.pid.nunique())
    r16 = agg.loc[agg.k == 16].iloc[0]
    r32 = agg.loc[agg.k == 32].iloc[0]
    r64 = agg.loc[agg.k == 64].iloc[0]
    remaining_gain_16_to_64 = float(r64.reconstruction_r2 - r16.reconstruction_r2)
    behavior_loss_16 = float(max(0, -r16.behavior_delta_r2_mean))

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1: minimal title.
    slide = prs.slides.add_slide(blank)
    _add_text(slide, "Track A", 0.75, 0.8, 3.0, 0.35, size=16, bold=True, color=COLORS["blue"])
    _add_text(slide, "Rate–distortion of low-dimensional\nLFP representations", 0.75, 1.45, 11.5, 1.6,
              size=34, bold=True)
    _add_text(slide, "Destriped 384-channel LFP → compact spatial basis → held-out reconstruction and behavior preservation",
              0.78, 3.35, 11.6, 0.7, size=18, color=COLORS["gray"])
    _add_text(slide, f"Pilot across {n_probes} successful public IBL probe{'s' if n_probes != 1 else ''} · {int(DURATION_S)} s per probe",
              0.78, 6.55, 11.0, 0.35, size=12, color=COLORS["gray"])

    # Slide 2: research question.
    slide = prs.slides.add_slide(blank)
    _add_title(slide, "One clean question: what do we lose as the representation gets smaller?", "Track A")
    y = 2.2
    labels = [
        ("384-channel\ndestriped LFP", COLORS["ink"]),
        ("k-dimensional\nspatial representation", COLORS["blue"]),
        ("reconstructed\nLFP", COLORS["orange"]),
        ("fixed ridge\nbehavior readout", COLORS["green"]),
    ]
    xs = [0.8, 3.8, 6.8, 9.8]
    for i, ((label, color), x) in enumerate(zip(labels, xs)):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.2), Inches(1.15))
        shape.fill.solid(); shape.fill.fore_color.rgb = _rgb("#F7F7F5")
        shape.line.color.rgb = _rgb(color); shape.line.width = Inches(0.018)
        _add_text(slide, label, x + 0.12, y + 0.22, 1.96, 0.72, size=17, bold=True, align="CENTER")
        if i < len(labels) - 1:
            _add_text(slide, "→", x + 2.35, y + 0.36, 0.45, 0.4, size=28, bold=True, color=COLORS["gray"], align="CENTER")
    _add_text(slide, "Change only k. Hold the reference signal, temporal split, feature extraction, and decoder protocol fixed.",
              1.0, 4.25, 11.2, 0.6, size=20, bold=True, align="CENTER")
    _add_footer(slide, "The representation is evaluated twice: native waveform reconstruction and downstream behavioral information.")

    # Slide 3: protocol.
    slide = prs.slides.add_slide(blank)
    _add_title(slide, "The protocol isolates representation quality", "Methods")
    _add_text(slide, "Fit", 0.85, 1.65, 1.0, 0.3, size=13, bold=True, color=COLORS["blue"])
    _add_text(slide, "Evaluate once", 10.25, 1.65, 2.0, 0.3, size=13, bold=True, color=COLORS["orange"])
    train = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(2.05), Inches(9.35), Inches(0.55))
    train.fill.solid(); train.fill.fore_color.rgb = _rgb("#DDEBFA"); train.line.fill.background()
    test = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.2), Inches(2.05), Inches(2.3), Inches(0.55))
    test.fill.solid(); test.fill.fore_color.rgb = _rgb("#FBE4DA"); test.line.fill.background()
    _add_text(slide, "first 80% of time", 3.5, 2.14, 3.0, 0.25, size=15, bold=True, align="CENTER")
    _add_text(slide, "final 20%", 10.35, 2.14, 2.0, 0.25, size=15, bold=True, align="CENTER")
    items = [
        ("Reference", "destriped LFP at 250 Hz"),
        ("Rate", "k = 2, 4, 8, 16, 32, 64 components"),
        ("Distortion 1", "held-out native waveform 1 − R²"),
        ("Distortion 2", "behavior R² loss versus raw LFP"),
        ("Behavior model", "same depth-group band-power features + RidgeCV"),
    ]
    for i, (head, body) in enumerate(items):
        yy = 3.05 + i * 0.68
        _add_text(slide, head, 1.0, yy, 2.1, 0.3, size=15, bold=True, color=COLORS["ink"])
        _add_text(slide, body, 3.05, yy, 8.9, 0.35, size=15, color=COLORS["gray"])
    _add_footer(slide, f"Pilot: {n_probes} successful probes; one-second spectral windows; windows crossing the 80/20 boundary are excluded.")

    # Slide 4: reconstruction.
    slide = prs.slides.add_slide(blank)
    _add_title(slide, f"At k = 16, median held-out waveform reconstruction R² is {r16.reconstruction_r2:.3f}", "Result 1")
    slide.shapes.add_picture(figure_paths["reconstruction"], Inches(0.65), Inches(1.35), width=Inches(12.0))
    _add_text(slide, f"Moving from 16 to 64 components adds {remaining_gain_16_to_64:.3f} median R²; the curve shows the marginal return directly rather than imposing a cutoff.",
              0.9, 6.65, 11.5, 0.4, size=13, color=COLORS["gray"])
    _add_footer(slide, "Thin lines: individual probes. Thick line: median. PCA and normalization are fit on training time only.")

    # Slide 5: traces.
    slide = prs.slides.add_slide(blank)
    _add_title(slide, "Increasing k restores channel-specific waveform detail", "Result 1 · example")
    slide.shapes.add_picture(figure_paths["traces"], Inches(0.65), Inches(1.35), width=Inches(12.0))
    _add_footer(slide, "Representative held-out traces from one probe. Vertical offsets separate reconstructions; they do not change waveform shape.")

    # Slide 6: behavior.
    slide = prs.slides.add_slide(blank)
    if np.isfinite(r16.behavior_delta_r2_mean):
        behavior_title = f"At k = 16, median mean behavior R² changes by {r16.behavior_delta_r2_mean:+.3f} versus raw LFP"
    else:
        behavior_title = "Behavioral preservation is evaluated separately from waveform reconstruction"
    _add_title(slide, behavior_title, "Result 2")
    slide.shapes.add_picture(figure_paths["behavior"], Inches(0.65), Inches(1.35), width=Inches(12.0))
    _add_text(slide, "Each curve compares a reconstructed signal with the raw destriped reference using the same downstream features and ridge-selection procedure.",
              0.9, 6.65, 11.5, 0.4, size=13, color=COLORS["gray"])
    _add_footer(slide, "R² may be negative on difficult held-out behaviors; the plotted quantity is paired within probe and behavior.")

    # Slide 7: two distortions.
    slide = prs.slides.add_slide(blank)
    _add_title(slide, "Track A has two distortions—and both belong on the curve", "Interpretation")
    slide.shapes.add_picture(figure_paths["combined"], Inches(0.65), Inches(1.35), width=Inches(12.0))
    _add_footer(slide, "Blue: loss of the signal itself. Orange: loss of information available to a fixed behavioral readout.")

    # Slide 8: interpretation and next step.
    slide = prs.slides.add_slide(blank)
    _add_title(slide, "What this pilot establishes—and what it does not", "Conclusion")
    columns = [
        ("Supported", [
            "A reproducible, train/test-separated rate–distortion protocol is now defined.",
            f"The curve is measured on {n_probes} public IBL probes using destriped data.",
            "Waveform and behavioral preservation are evaluated under matched conditions.",
        ], COLORS["green"]),
        ("Not yet supported", [
            "A universal optimal k: the mentor specified no hard cutoff.",
            "Generalization across the full brain-wide dataset from this small pilot.",
            "Superiority of PCA over other low-dimensional representation families.",
        ], COLORS["orange"]),
    ]
    for idx, (head, bullets, color) in enumerate(columns):
        x = 0.8 + idx * 6.15
        _add_text(slide, head, x, 1.65, 5.5, 0.4, size=20, bold=True, color=color)
        for j, b in enumerate(bullets):
            _add_text(slide, "• " + b, x, 2.3 + j * 1.05, 5.45, 0.78, size=16, color=COLORS["ink"])
    next_text = "Next: repeat on the larger probe cohort, then compare PCA with one alternative representation using the identical curve and decoder."
    _add_text(slide, next_text, 0.9, 6.18, 11.5, 0.58, size=18, bold=True, align="CENTER")
    failure_note = ""
    if failures:
        failure_note = f" {len(failures)} preselected probe(s) failed data access and are listed in the run manifest."
    _add_footer(slide, "This is a pilot result, not a final cohort-level conclusion." + failure_note)

    out = OUT / "Track_A_Rate_Distortion.pptx"
    prs.save(out)
    return out


def flatten_results(rows: Sequence[ProbeResult]) -> pd.DataFrame:
    records = []
    for r in rows:
        d = asdict(r)
        d["behavior_scores"] = r.behavior_scores
        d["raw_behavior_scores"] = r.raw_behavior_scores
        records.append(d)
    return pd.DataFrame(records)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    FIG.mkdir(parents=True, exist_ok=True)
    started = time.time()
    one = ONE(
        base_url="https://openalyx.internationalbrainlab.org",
        password="international",
        silent=True,
    )

    all_rows: List[ProbeResult] = []
    examples: List[Dict[str, object]] = []
    failures: List[Dict[str, str]] = []
    for spec in PROBES:
        try:
            rows, example = run_one_probe(one, spec)
            all_rows.extend(rows)
            examples.append(example)
        except Exception as exc:
            traceback.print_exc()
            failures.append({
                "pid": str(spec["pid"]),
                "label": str(spec["label"]),
                "error": f"{type(exc).__name__}: {exc}",
            })

    if not all_rows:
        raise RuntimeError("Track A failed: no preselected probes completed")

    df = flatten_results(all_rows)
    # JSON dict columns are kept in the detailed JSON; CSV expands behavior scores below.
    csv_rows = []
    for r in all_rows:
        base = {
            k: v for k, v in asdict(r).items()
            if k not in ("behavior_scores", "raw_behavior_scores")
        }
        for name, value in r.behavior_scores.items():
            base[f"behavior_r2_{name}"] = value
        for name, value in r.raw_behavior_scores.items():
            base[f"raw_behavior_r2_{name}"] = value
        csv_rows.append(base)
    pd.DataFrame(csv_rows).to_csv(OUT / "track_a_per_probe.csv", index=False)
    aggregate_results(df).to_csv(OUT / "track_a_summary.csv", index=False)

    figure_paths = make_figures(df, examples)
    deck = make_deck(df, figure_paths, failures)
    manifest = {
        "status": "completed",
        "started_utc": pd.Timestamp.utcnow().isoformat(),
        "runtime_s": time.time() - started,
        "protocol": {
            "reference": "destriped 384-channel LFP at 250 Hz",
            "method": "spatial PCA fit on first 80% of each probe",
            "dimensions": list(DIMS),
            "duration_s": DURATION_S,
            "train_fraction": TRAIN_FRACTION,
            "reconstruction_metric": "pooled held-out R2 against destriped data",
            "behavior_metric": "per-feature held-out prediction R2 from fixed band-power features and RidgeCV",
        },
        "successful_probes": sorted(df.pid.unique().tolist()),
        "failures": failures,
        "outputs": {
            "deck": str(deck),
            "per_probe_csv": str(OUT / "track_a_per_probe.csv"),
            "summary_csv": str(OUT / "track_a_summary.csv"),
            "figures": figure_paths,
        },
    }
    (OUT / "run_manifest.json").write_text(json.dumps(manifest, indent=2))
    print(json.dumps(manifest, indent=2), flush=True)


if __name__ == "__main__":
    main()
